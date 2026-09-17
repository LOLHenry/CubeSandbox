#!/usr/bin/env python3
"""Slide: MAI-UI Fig.5 topology with callouts matching the two form bullets."""

from pathlib import Path

import urllib.request
from PIL import Image, ImageDraw, ImageFont

ROOT = Path("/workspace/docs/assets")
ART = Path("/opt/cursor/artifacts")
FONT = "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"
FIG5_URL = "https://arxiv.org/html/2512.22047v1/rl_train_schema2.png"

NAVY = (22, 42, 90)
BLUE = (47, 111, 212)
ORANGE = (201, 108, 22)
RED = (196, 48, 48)
WHITE = (255, 255, 255)
TEXT = (28, 34, 48)
MUTED = (90, 98, 114)
ORANGE_BG = (255, 243, 224)
BLUE_BG = (232, 242, 255)
CREAM = (255, 252, 246)


def font(size):
    return ImageFont.truetype(FONT, size)


def round_rect(draw, xy, r, fill, outline, width=3):
    kw = {"outline": outline, "width": width}
    if fill is not None:
        kw["fill"] = fill
    draw.rounded_rectangle(xy, radius=r, **kw)


def badge(draw, xy, n, fill):
    x, y, s = xy
    draw.ellipse((x, y, x + s, y + s), fill=fill, outline=WHITE, width=3)
    f = font(int(s * 0.55))
    t = str(n)
    bb = draw.textbbox((0, 0), t, font=f)
    draw.text(
        (x + (s - (bb[2] - bb[0])) / 2, y + (s - (bb[3] - bb[1])) / 2 - 2),
        t,
        font=f,
        fill=WHITE,
    )


def fetch_fig5():
    dest = ART / "maiui-fig5-original.png"
    ART.mkdir(parents=True, exist_ok=True)
    if not dest.exists():
        urllib.request.urlretrieve(FIG5_URL, dest)
    return Image.open(dest).convert("RGB")


def annotate_fig5(src):
    """Keep Fig.5 pixels; add ① on Screenshots, ② on CPU env stack."""
    img = src.copy()
    w, h = img.size
    d = ImageDraw.Draw(img, "RGBA")

    # ① Screenshots path (upper-right dashed arrows)
    x1, y1, x2, y2 = int(w * 0.78), int(h * 0.18), int(w * 0.97), int(h * 0.42)
    d.rounded_rectangle((x1, y1, x2, y2), radius=14, outline=ORANGE + (255,), width=8)
    badge(d, (x1 - 28, y1 - 18, 56), 1, ORANGE)

    # ② CPU Worker env stack
    x1, y1, x2, y2 = int(w * 0.70), int(h * 0.40), int(w * 0.97), int(h * 0.92)
    d.rounded_rectangle((x1, y1, x2, y2), radius=14, outline=BLUE + (255,), width=8)
    badge(d, (x1 - 28, y2 - 50, 56), 2, BLUE)

    # ① also ticks rollout wall-clock (left trainer)
    x1, y1, x2, y2 = int(w * 0.018), int(h * 0.28), int(w * 0.195), int(h * 0.52)
    d.rounded_rectangle((x1, y1, x2, y2), radius=12, outline=ORANGE + (180,), width=5)

    cap = Image.new("RGB", (w, h + 88), WHITE)
    cap.paste(img, (0, 0))
    c = ImageDraw.Draw(cap)
    c.text((24, h + 12), "① 截图边：Screenshots → Agent 等稳定帧。出帧慢，拉长采轨迹墙钟。", font=font(26), fill=ORANGE)
    c.text((24, h + 48), "② 装箱边：鲲鹏沙箱节点上路数。决策在昇腾上算，沙箱仍出帧，占满才开不满。论文 GPU Worker = 我侧昇腾节点。", font=font(26), fill=BLUE)
    return cap


def slide(annotated):
    """16:9 胶片：左图右文。"""
    W, H = 1920, 1080
    img = Image.new("RGB", (W, H), WHITE)
    d = ImageDraw.Draw(img)

    d.rectangle((0, 0, W, 88), fill=NAVY)
    d.text((36, 22), "GUI Agent 依据渲染稳定后的截图决策 —— 渲染打在训练环的两处", font=font(32), fill=WHITE)
    d.text((36, 58), "左图结构对齐 MAI-UI Fig.5。论文写 GPU Worker，我侧改称昇腾节点 / 鲲鹏沙箱节点。①② 与右边两段对应。", font=font(18), fill=(190, 205, 230))

    # left: fit annotated fig
    aw, ah = annotated.size
    box_w, box_h = 1120, 940
    scale = min(box_w / aw, box_h / ah)
    nw, nh = int(aw * scale), int(ah * scale)
    fig = annotated.resize((nw, nh), Image.Resampling.LANCZOS)
    img.paste(fig, (24, 110 + (box_h - nh) // 2))

    # right cards
    rx = 1168
    round_rect(d, (rx, 110, 1888, 560), 16, ORANGE_BG, ORANGE, 4)
    badge(d, (rx + 24, 128, 44), 1, ORANGE)
    d.text((rx + 80, 132), "截图阶段 · 打在 Screenshots", font=font(26), fill=ORANGE)
    t1 = (
        "复杂界面软渲染热点可超过 50%，帧率 < 10fps。\n"
        "出帧慢，Agent 拿到稳定截图变晚，\n"
        "单步等待变长，一轮 Rollout 增加 5%~15%*"
    )
    d.multiline_text((rx + 28, 196), t1, font=font(24), fill=TEXT, spacing=10)
    d.text((rx + 28, 430), "图中落点：右上 Screenshots 虚线，\n并传导到左侧 Multi-turn Online Rollout。", font=font(20), fill=MUTED)

    round_rect(d, (rx, 584, 1888, 980), 16, BLUE_BG, BLUE, 4)
    badge(d, (rx + 24, 602, 44), 2, BLUE)
    d.text((rx + 80, 606), "非截图阶段 · 打在鲲鹏沙箱路数", font=font(26), fill=BLUE)
    t2 = (
        "决策期不需要新帧，软渲染仍在出帧，\n"
        "较多占用 CPU，单机并发沙箱密度减少。\n"
        "可优化空间 20%+†"
    )
    d.multiline_text((rx + 28, 670), t2, font=font(24), fill=TEXT, spacing=10)
    d.text((rx + 28, 850), "图中落点：右侧 Env1…N 堆叠。\nFig.5 本身不画「停刷」，要靠 ② 圈出这叠沙箱。", font=font(20), fill=MUTED)

    d.text(
        (36, 1044),
        "* 中小模型、单步约 2–3s，截图多等 0.1–0.3s。† 全程 60→20fps 实验上限，不是按需已兑现；加路还要整机超卖且截图能交错。",
        font=font(16),
        fill=MUTED,
    )
    return img


def schematic():
    """Original-drawn topology (do not copy paper pixels into git)."""
    W, H = 1600, 780
    img = Image.new("RGB", (W, H), CREAM)
    d = ImageDraw.Draw(img)
    d.text((36, 20), "结构对齐 MAI-UI Fig.5（自绘，便于胶片标注）", font=font(26), fill=NAVY)

    round_rect(d, (36, 80, 300, 700), 16, WHITE, NAVY, 3)
    d.text((56, 100), "训练器", font=font(22), fill=NAVY)
    round_rect(d, (56, 160, 280, 300), 10, ORANGE_BG, ORANGE, 3)
    d.text((72, 188), "在线采轨迹", font=font(22), fill=ORANGE)
    d.text((72, 228), "Multi-turn Rollout", font=font(16), fill=MUTED)
    round_rect(d, (56, 340, 280, 460), 10, (240, 240, 246), MUTED, 2)
    d.text((72, 380), "策略更新", font=font(22), fill=NAVY)
    d.text((72, 420), "Policy Update", font=font(16), fill=MUTED)
    badge(d, (250, 150, 40), 1, ORANGE)

    round_rect(d, (340, 80, 980, 700), 16, WHITE, BLUE, 3)
    d.text((360, 100), "昇腾节点  ·  模型侧 Agent 循环", font=font(22), fill=BLUE)
    for i, name in enumerate(("任务 1", "任务 2", "任务 N")):
        y = 180 + i * 150
        round_rect(d, (370, y, 520, y + 70), 8, (255, 236, 214), ORANGE, 2)
        d.text((390, y + 20), name, font=font(20), fill=ORANGE)
        round_rect(d, (560, y, 780, y + 70), 8, BLUE_BG, BLUE, 2)
        d.text((580, y + 20), "Agent Loop", font=font(20), fill=BLUE)
        round_rect(d, (820, y, 950, y + 70), 8, (232, 248, 232), (46, 125, 50), 2)
        d.text((838, y + 20), "轨迹", font=font(20), fill=(46, 125, 50))

    round_rect(d, (1020, 80, 1564, 700), 16, WHITE, (46, 125, 50), 3)
    d.text((1040, 96), "鲲鹏节点  ·  Android 沙箱", font=font(22), fill=(46, 125, 50))
    d.text((1040, 132), "Actions ↓          Screenshots ↑", font=font(16), fill=ORANGE)
    round_rect(d, (1288, 124, 1548, 168), 8, ORANGE_BG, ORANGE, 3)
    d.text((1300, 132), "① 截图边", font=font(18), fill=ORANGE)
    for i, name in enumerate(("Env 1", "Env 2", "Env N")):
        y = 196 + i * 140
        round_rect(d, (1100, y, 1480, y + 90), 10, (232, 248, 232), (46, 125, 50), 3)
        d.text((1180, y + 28), f"在线沙箱 {name}", font=font(22), fill=NAVY)
    round_rect(d, (1080, 184, 1508, 620), 12, None, BLUE, 4)
    badge(d, (1488, 128, 44), 1, ORANGE)
    badge(d, (1488, 560, 44), 2, BLUE)
    d.text((1040, 648), "① 等稳定截图    ② 决策期多路仍出帧", font=font(20), fill=TEXT)
    return img


def write_pptx(schematic_path):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR
    from pptx.util import Emu, Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    def box(l, t, w, h, fill, line):
        sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor(*fill)
        sh.line.color.rgb = RGBColor(*line)
        sh.line.width = Pt(1.75)
        return sh

    def text(l, t, w, h, content, size, color, bold=False):
        tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = content
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor(*color)
        run.font.name = "微软雅黑"
        return tb

    head = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.72))
    head.fill.solid()
    head.fill.fore_color.rgb = RGBColor(*NAVY)
    head.line.fill.background()
    text(0.28, 0.12, 12.7, 0.5, "GUI Agent 依据渲染稳定后的截图进行决策", 22, WHITE, True)

    slide.shapes.add_picture(str(schematic_path), Inches(0.18), Inches(0.88), Inches(8.15), Inches(5.95))

    c1 = box(8.5, 0.88, 4.55, 2.85, ORANGE_BG, ORANGE)
    tf = c1.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "① 截图阶段 · Screenshots\n复杂界面软渲染热点可超过 50%，帧率 < 10fps。出帧慢，Agent 拿到稳定截图变晚，单步等待变长，一轮 Rollout 增加 5%~15%*"
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(*TEXT)
    r.font.name = "微软雅黑"

    c2 = box(8.5, 3.88, 4.55, 2.55, BLUE_BG, BLUE)
    tf = c2.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "② 非截图阶段 · 鲲鹏沙箱路数\n决策期不需要新帧，软渲染仍在出帧，较多占用 CPU，单机并发沙箱密度减少。可优化空间 20%+†"
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(*TEXT)
    r.font.name = "微软雅黑"

    text(
        0.28,
        6.95,
        12.7,
        0.4,
        "* 中小模型、单步约 2–3s。† 全程 60→20fps 实验上限，不是按需已兑现。左图为结构草图：昇腾节点跑模型，鲲鹏节点跑沙箱。论文 Fig.5 写 GPU Worker，口播改称昇腾节点。",
        11,
        MUTED,
    )

    out = ROOT / "maiui-fig5-slide.pptx"
    prs.save(out)
    prs.save(ART / "maiui-fig5-slide.pptx")
    return out


def main():
    ART.mkdir(parents=True, exist_ok=True)
    ROOT.mkdir(parents=True, exist_ok=True)
    raw = fetch_fig5()
    ann = annotate_fig5(raw)
    ann.save(ART / "maiui-fig5-annotated.png")
    s = slide(ann)
    s.save(ART / "maiui-fig5-slide.png")
    sch = schematic()
    sch.save(ROOT / "maiui-fig5-schematic.png")
    sch.save(ART / "maiui-fig5-schematic.png")
    pptx = write_pptx(ROOT / "maiui-fig5-schematic.png")
    print("wrote", ART / "maiui-fig5-annotated.png")
    print("wrote", ART / "maiui-fig5-slide.png")
    print("wrote", ROOT / "maiui-fig5-schematic.png")
    print("wrote", pptx)


if __name__ == "__main__":
    main()
